"""Extract plain text from uploaded documents (PDF, DOCX, TXT)."""

from __future__ import annotations

import io

from app.core.errors import AppError


def detect_kind(filename: str, mime: str | None) -> str:
    name = filename.lower()
    if name.endswith(".pdf") or (mime or "").endswith("pdf"):
        return "pdf"
    if name.endswith(".docx") or "word" in (mime or ""):
        return "docx"
    if name.endswith(".txt") or (mime or "").startswith("text/"):
        return "txt"
    raise AppError("Unsupported file type (use PDF, DOCX, or TXT)", code="unsupported_file")


def extract_text(data: bytes, filename: str, mime: str | None = None) -> str:
    kind = detect_kind(filename, mime)
    if kind == "txt":
        return data.decode("utf-8", errors="replace")
    if kind == "pdf":
        return _extract_pdf(data)
    return _extract_docx(data)


def _extract_pdf(data: bytes) -> str:
    # pdfplumber gives better layout-aware text; fall back to pypdf.
    try:
        import pdfplumber

        parts: list[str] = []
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages:
                parts.append(page.extract_text() or "")
        text = "\n".join(parts).strip()
        if text:
            return text
    except Exception:
        pass

    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n".join((page.extract_text() or "") for page in reader.pages).strip()


def _extract_docx(data: bytes) -> str:
    from docx import Document as DocxDocument

    doc = DocxDocument(io.BytesIO(data))
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(ln for ln in lines if ln.strip()).strip()
